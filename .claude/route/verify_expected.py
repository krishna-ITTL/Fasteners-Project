"""Independent implementation of the VERIFICATION checks.

Ground truth for testing RunVerification. Deliberately written from the spec
rather than from the VBA, so agreement between the two means something.

Usage:  python verify_expected.py <workbook.xlsm> [--json out.json]
"""
import sys, re, json, collections
import openpyxl

TH_FIRST, TH_LAST = 15, 250
GROUP_MIN = 5


def num(v):
    return v if isinstance(v, (int, float)) else 0


def load_standard(sd):
    std = {}
    for r in range(222, 228):
        ref = str(sd.cell(r, 11).value or "").strip()
        if not ref:
            continue
        std[ref] = dict(name=sd.cell(r, 5).value, bolt=sd.cell(r, 7).value,
                        nut=sd.cell(r, 8).value, pw=sd.cell(r, 9).value,
                        sw=sd.cell(r, 10).value)
    return std


def finding(row, joint, loc, des, item, sev, txt, gkey):
    """gkey: what makes two findings 'the same kind' for grouping."""
    return dict(row=row, joint=joint, loc=loc, des=des, item=item,
                sev=sev, txt=txt, gkey=gkey)


def check(path, rulings=None):
    """rulings: {(joint, item): accepted_value} overriding the standard."""
    rulings = rulings or {}
    wb = openpyxl.load_workbook(path, data_only=True)
    th, sd = wb["TANK HARDWARE"], wb["STD.DATA"]
    std = load_standard(sd)
    out = []

    for r in range(TH_FIRST, TH_LAST + 1):
        M = th.cell(r, 13).value
        if M in (None, ""):
            continue
        joint = str(th.cell(r, 4).value or "").strip()
        loc = str(th.cell(r, 2).value or "")
        des = str(M)
        N, O = num(th.cell(r, 14).value), num(th.cell(r, 15).value)
        Q, R = num(th.cell(r, 17).value), num(th.cell(r, 18).value)

        if joint not in std:
            out.append(finding(r, joint or "(blank)", loc, des, "-", "UNVERIFIABLE",
                               "joint type not in STD.DATA",
                               (joint, "-", "UNVERIFIABLE", None)))
            continue

        s = std[joint]
        exp = lambda item, v: rulings.get((joint, item), v)

        if s["bolt"] > 0:
            if N == 0:
                out.append(finding(r, joint, loc, des, "BOLT/STUD", "INFO",
                                   "standard expects a fastener but quantity is 0",
                                   (joint, "BOLT/STUD-ZERO", "INFO", None)))
                continue
            for item, act, want in (("NUT", O / N, exp("NUT", s["nut"])),
                                    ("P.WASHER", Q / N, exp("P.WASHER", s["pw"])),
                                    ("S.WASHER", R / N, exp("S.WASHER", s["sw"]))):
                act = round(act, 4)
                if act != want:
                    out.append(finding(r, joint, loc, des, item, "ERROR",
                                       "%g per bolt vs standard %g" % (act, want),
                                       (joint, item, "ERROR", act)))
        else:
            # Welded pad: a per-bolt ratio is undefined, compare against nuts.
            if N != 0:
                out.append(finding(r, joint, loc, des, "BOLT/STUD", "ERROR",
                                   "standard buys no fastener, qty is %g" % N,
                                   (joint, "BOLT/STUD-PRESENT", "ERROR", None)))
            if O > 0:
                for item, act, want in (("P.WASHER", Q / O, exp("P.WASHER", s["pw"] / s["nut"])),
                                        ("S.WASHER", R / O, exp("S.WASHER", s["sw"] / s["nut"]))):
                    act = round(act, 4)
                    if act != want:
                        out.append(finding(r, joint, loc, des, item, "ERROR",
                                           "%g per nut vs standard %g" % (act, want),
                                           (joint, item, "ERROR", act)))

    # --- check C ---
    for r in range(TH_FIRST, TH_LAST + 1):
        M = th.cell(r, 13).value
        if M in (None, ""):
            continue
        K, L, B = th.cell(r, 11).value, th.cell(r, 12).value, th.cell(r, 2).value
        des = str(M)
        if isinstance(K, (int, float)) and isinstance(L, (int, float)) and L < K:
            out.append(finding(r, "", str(B or ""), des, "LENGTH", "ERROR",
                               "chosen length %g shorter than required %g" % (L, K),
                               ("", "LENGTH", "ERROR", None)))
        if B and M:
            w = re.match(r"M(\d+)\s*(?:-STUD)?[xX](\d+)", des, re.I)
            for a, b in re.findall(r"M(\d+)\s*[xX]\s*(\d+)", str(B)):
                if w and (a, b) != (w.group(1), w.group(2)):
                    out.append(finding(r, "", str(B)[:40], des, "TEXT", "INFO",
                                       "description says M%sx%s" % (a, b),
                                       ("", "TEXT", "INFO", None)))
    return out


def group(findings):
    buckets = collections.defaultdict(list)
    for f in findings:
        buckets[f["gkey"]].append(f)
    grouped, individual = [], []
    for k, v in buckets.items():
        if len(v) >= GROUP_MIN:
            grouped.append((k, v))
        else:
            individual.extend(v)
    grouped.sort(key=lambda kv: -len(kv[1]))
    individual.sort(key=lambda f: (f["row"], f["item"]))
    return grouped, individual


if __name__ == "__main__":
    path = sys.argv[1]
    f = check(path)
    grouped, individual = group(f)
    print("raw findings: %d  ->  %d grouped lines + %d individual"
          % (len(f), len(grouped), len(individual)))
    print()
    print("GROUPED")
    for k, v in grouped:
        print("   %-9s %-10s %-12s %-34s %d rows"
              % (v[0]["joint"], v[0]["item"], v[0]["sev"], v[0]["txt"], len(v)))
    print()
    print("INDIVIDUAL")
    for x in individual:
        print("   r%-4d %-9s %-26s %-10s %-12s %s"
              % (x["row"], x["joint"], x["loc"][:26], x["item"], x["sev"], x["txt"]))
    if "--json" in sys.argv:
        out = sys.argv[sys.argv.index("--json") + 1]
        json.dump({"grouped": [[v[0]["joint"], v[0]["item"], v[0]["sev"], v[0]["txt"], len(v)]
                               for k, v in grouped],
                   "individual": [[x["row"], x["joint"], x["item"], x["sev"], x["txt"]]
                                  for x in individual]},
                  open(out, "w"), indent=1)
        print("\nsaved -> " + out)
