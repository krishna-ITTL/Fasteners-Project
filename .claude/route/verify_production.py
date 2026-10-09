"""Independent check of PRODUCTION SHEET against PRINT SHEET and FINAL SHEET.

Usage: python verify_production.py <workbook.xlsm>   (read-only, cached values)
Exits non-zero on any mismatch.
"""
import re
import sys
import openpyxl

path = sys.argv[1]
wb = openpyxl.load_workbook(path, data_only=True)
ps, pd, fs = wb["PRINT SHEET"], wb["PRODUCTION SHEET"], wb["FINAL SHEET"]
pat = re.compile(r"^\s*(M\d+)\s*(-STUD)?\s*x\s*(\d+)\s*-\s*(.+?)\s*$", re.I)

# --- expected, rebuilt from PRINT SHEET
top = {}
for m in ps.merged_cells.ranges:
    if m.min_col <= 4 <= m.max_col:
        for r in range(m.min_row, m.max_row + 1):
            top[r] = m.min_row
last = max(245, max(r for r in range(1, ps.max_row + 1) if ps.cell(r, 6).value is not None))
groups = {}
for r in range(9, last + 1):
    f = ps.cell(r, 6).value
    if f is None or str(f).strip() == "":
        continue
    q = [ps.cell(r, c).value for c in range(7, 12)]
    q = [v if isinstance(v, (int, float)) else 0 for v in q]
    if not any(q):
        continue
    assert pat.match(str(f)), f"bad designation row {r}: {f}"
    loc = str(ps.cell(top.get(r, r), 4).value or "").strip()
    m = pat.match(str(f))
    key = m.group(1).upper() + ("-STUD" if m.group(2) else "") + "X" + str(int(m.group(3))) + "-" + m.group(4).strip()
    groups.setdefault(key, []).append([loc] + q + [ps.cell(r, 12).value])

def sort_code(k):
    m = pat.match(k)
    return (1 if m.group(2) else 0, int(m.group(1)[1:]), int(m.group(3)), m.group(4))

expected = [(k, groups[k]) for k in sorted(groups, key=sort_code)]

# --- actual, read from PRODUCTION SHEET
actual, r, maxr = [], 7, pd.max_row
while r <= maxr:
    title = pd.cell(r, 4).value
    if title is None:
        r += 1
        continue
    head = [pd.cell(r + 1, c).value for c in range(4, 12)]
    assert head[0] == "JOINT LOCATION", f"row {r+1} header {head}"
    assert head[2] == ("STUD" if "-STUD" in title else "BOLT"), f"row {r+1} {head}"
    rows, rr = [], r + 2
    while pd.cell(rr, 4).value != "TOTAL":
        assert rr < r + 500, f"block at row {r} has no TOTAL row"
        rows.append([pd.cell(rr, 4).value] + [pd.cell(rr, c).value for c in range(6, 12)])
        rr += 1
    tot = [pd.cell(rr, c).value for c in range(6, 11)]
    want = [sum(x[i] for x in rows) for i in range(1, 6)]
    assert tot == want, f"TOTAL row {rr}: {tot} != {want}"
    actual.append((title, rows))
    last_total = rr
    r = rr + 1

errors = 0
if [k for k, _ in expected] != [k for k, _ in actual]:
    errors += 1
    print("BLOCK ORDER/SET differs\n exp:", [k for k, _ in expected], "\n act:", [k for k, _ in actual])
for (ek, er), (ak, ar) in zip(expected, actual):
    if er != ar:
        errors += 1
        print("BLOCK", ek, "\n exp:", er, "\n act:", ar)

# --- grand totals must match FINAL SHEET
# Row ranges = FS_B1/B2/B3 constants in FinalSheetBuilder; update both together.
def fs_sum(first, last, col):
    return sum(v for v in (fs.cell(r, col).value for r in range(first, last + 1)) if isinstance(v, (int, float)))

ln = sum(x[5] for _, rows in actual for x in rows)
bolts = sum(x[1] for k, rows in actual if "-STUD" not in k for x in rows)
studs = sum(x[1] for k, rows in actual if "-STUD" in k for x in rows)
nuts = sum(x[2] for _, rows in actual for x in rows)
pw = sum(x[3] for _, rows in actual for x in rows)
sw = sum(x[4] for _, rows in actual for x in rows)
fs_tot = dict(bolts=fs_sum(37, 96, 7), studs=fs_sum(99, 125, 7),
              nuts=fs_sum(11, 34, 7), pw=fs_sum(11, 34, 8), sw=fs_sum(11, 34, 9), ln=fs_sum(11, 34, 10))
pd_tot = dict(bolts=bolts, studs=studs, nuts=nuts, pw=pw, sw=sw, ln=ln)
print("PRODUCTION totals:", pd_tot)
print("FINAL      totals:", fs_tot)
if pd_tot != fs_tot:
    errors += 1
    print("TOTALS DIFFER")

print("print area:", pd.print_area, "| title rows:", pd.print_title_rows)
want_area = f"'PRODUCTION SHEET'!$D$2:$K${last_total}"
if pd.print_area != want_area or pd.print_title_rows != "$2:$6":
    errors += 1
    print("PRINT SETUP wrong, expected", want_area)
print(f"{len(actual)} blocks, {sum(len(x) for _, x in actual)} rows, errors={errors}")
sys.exit(1 if errors else 0)
